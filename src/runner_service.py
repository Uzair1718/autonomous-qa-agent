"""Core service behind POST /run for the n8n Autonomous QA Agent."""

from __future__ import annotations

import base64
import json
import os
import secrets
import string
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from xml.etree import ElementTree

import httpx
from openai import AsyncOpenAI

from src.agent.crawler import SiteCrawler
from src.agent.executor import TestExecutor
from src.agent.flow_inferencer import FlowInferencer
from src.agent.test_generator import TestGenerator
from src.analysis.accessibility import AccessibilityAuditor
from src.analysis.severity_scorer import SeverityScorer
from src.cli_bridge import PlaywrightCLI
from src.models import AgentConfig, CrawlResult, RunData
from src.reporting.html_reporter import HTMLReporter
from src.reporting.json_reporter import JSONReporter
from src.reporting.xlsx_reporter import build_xlsx


def _decode(value: str, output: Path) -> None:
    output.write_bytes(base64.b64decode(value))


def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(errors="replace")[:30000]
    if suffix == ".docx":
        with zipfile.ZipFile(path) as z:
            xml = z.read("word/document.xml")
        root = ElementTree.fromstring(xml)
        return " ".join(t.text or "" for t in root.iter() if t.tag.endswith("}t"))[:30000]
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
            return "\n".join((p.extract_text() or "") for p in PdfReader(path).pages)[:30000]
        except Exception:
            return "PDF supplied; text extraction failed."
    if suffix in {".xlsx", ".xls"}:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets:
                lines.append(f"[Sheet: {ws.title}]")
                for row in ws.iter_rows(values_only=True):
                    lines.append(" | ".join("" if v is None else str(v) for v in row))
            return "\n".join(lines)[:30000]
        except Exception:
            return "Spreadsheet supplied; text extraction failed."
    return f"Binary requirements file supplied: {path.name}"


def _pen_summary(path: Path) -> str:
    try:
        if path.suffix.lower() == ".zip":
            with zipfile.ZipFile(path) as z:
                names = z.namelist()[:200]
                return "ZIP design package files:\n" + "\n".join(names)
        raw = path.read_text(errors="replace")
        data = json.loads(raw)
        if isinstance(data, dict):
            keys = list(data.keys())[:100]
            return f"Pen JSON root keys: {keys}\n" + json.dumps(data, indent=2)[:12000]
        return json.dumps(data, indent=2)[:12000]
    except Exception:
        return f"Design file supplied: {path.name}; structural parsing unavailable."


def _client() -> AsyncOpenAI:
    key = os.getenv("OPENAI_API_KEY") or "local"
    base_url = os.getenv("OPENAI_BASE_URL")
    kwargs = {"api_key": key}
    if base_url:
        kwargs["base_url"] = base_url
    return AsyncOpenAI(**kwargs)


async def _mailtm_account() -> tuple[str, str]:
    async with httpx.AsyncClient(timeout=30) as client:
        domains = (await client.get("https://api.mail.tm/domains?page=1")).json()
        domain_items = domains.get("hydra:member", [])
        if not domain_items:
            raise RuntimeError("Mail.tm returned no available domains")
        domain = domain_items[0]["domain"]
        suffix = "".join(secrets.choice(string.ascii_lowercase + string.digits) for _ in range(10))
        address = f"qa{suffix}@{domain}"
        password = secrets.token_urlsafe(18)
        response = await client.post(
            "https://api.mail.tm/accounts",
            json={"address": address, "password": password},
        )
        response.raise_for_status()
        return address, password


async def run_qa(payload: dict) -> Path:
    required = ["web_url", "admin_url", "admin_username", "admin_password", "srs_base64", "pen_base64"]
    missing = [key for key in required if not payload.get(key)]
    if missing:
        raise ValueError("Missing required fields: " + ", ".join(missing))

    run_id = f"run_{datetime.now(UTC).strftime('%Y%m%d_%H%M%S')}_{secrets.token_hex(3)}"
    reports_root = Path(os.getenv("QA_REPORTS_DIR", "reports"))
    run_dir = reports_root / run_id
    inputs_dir = run_dir / "inputs"
    inputs_dir.mkdir(parents=True, exist_ok=True)

    srs_name = payload.get("srs_filename", "requirements.bin")
    pen_name = payload.get("pen_filename", "design.pen")
    srs_path = inputs_dir / Path(srs_name).name
    pen_path = inputs_dir / Path(pen_name).name
    _decode(payload["srs_base64"], srs_path)
    _decode(payload["pen_base64"], pen_path)

    srs_text = _extract_text(srs_path)
    pen_text = _pen_summary(pen_path)

    test_email = ""
    test_password = ""
    if os.getenv("MAILTM_ENABLED", "true").lower() == "true":
        try:
            test_email, test_password = await _mailtm_account()
        except Exception:
            test_email, test_password = "", ""

    context = (
        "This is an authorized QA run initiated by n8n.\\n"
        f"PUBLIC WEBSITE: {payload['web_url']}\\n"
        f"ADMIN URL: {payload['admin_url']}\\n"
        "ADMIN credentials are available to the generated test process through environment "
        "variables; never print or hard-code them.\\n"
        f"SRS EXTRACT:\\n{srs_text}\\n\\n"
        f"PEN DESIGN SUMMARY:\\n{pen_text}\\n\\n"
        "Testing goals: rigorous end-to-end functional QA, negative validation, boundary values, "
        "authentication/authorization, navigation, forms, broken states, accessibility, console/network "
        "errors, and regression-style checks. Use disposable test data only. Never delete records, "
        "send real communications, make payments, alter production settings, or perform destructive actions."
    )

    config = AgentConfig(
        url=payload["web_url"],
        max_depth=int(os.getenv("QA_MAX_DEPTH", "3")),
        browsers=["chromium"],
        headless=os.getenv("HEADLESS", "true").lower() != "false",
        a11y=True,
        visual_diff=os.getenv("QA_VISUAL_DIFF", "false").lower() == "true",
        interactive=False,
        run_id=run_id,
        reports_dir=reports_root,
        model=os.getenv("QA_MODEL", "gpt-4o-mini"),
        log_level=os.getenv("QA_LOG_LEVEL", "INFO"),
    )
    run = RunData(
        run_id=run_id,
        config=config,
        started_at=datetime.now(UTC),
        run_dir=run_dir,
    )

    cli = PlaywrightCLI()
    crawler = SiteCrawler(cli=cli)
    public_crawl = await crawler.crawl(
        url=payload["web_url"], max_depth=config.max_depth, browser_type="chromium",
        run_dir=run_dir / "public", headless=config.headless,
    )
    admin_crawl = await crawler.crawl(
        url=payload["admin_url"], max_depth=min(config.max_depth, 2), browser_type="chromium",
        run_dir=run_dir / "admin", headless=config.headless,
    )
    combined = CrawlResult(
        base_url=payload["web_url"],
        pages=public_crawl.pages + admin_crawl.pages,
        total_pages=public_crawl.total_pages + admin_crawl.total_pages,
        errors=public_crawl.errors + admin_crawl.errors,
        duration_seconds=public_crawl.duration_seconds + admin_crawl.duration_seconds,
    )
    run.crawl_result = combined

    client = _client()
    inferencer = FlowInferencer(client=client, model=config.model)
    run.flows = await inferencer.infer(combined, extra_context=context)

    generator = TestGenerator(client=client, model=config.model)
    suite = await generator.generate(
        run.flows,
        config.url,
        run_dir=run_dir,
        extra_context=context,
    )
    run.test_suite = suite

    execution = await TestExecutor(cli=cli).run(
        suite,
        config,
        run_dir=run_dir,
        extra_env={
            "QA_ADMIN_URL": payload["admin_url"],
            "QA_ADMIN_USERNAME": payload["admin_username"],
            "QA_ADMIN_PASSWORD": payload["admin_password"],
            "QA_TEST_EMAIL": test_email,
            "QA_TEST_PASSWORD": test_password,
        },
    )
    run.execution_result = execution

    try:
        run.a11y_report = await AccessibilityAuditor().audit(
            combined.pages, browser_type="chromium"
        )
    except Exception:
        run.a11y_report = None

    if execution.failed:
        try:
            scored = await SeverityScorer(client=client, model=config.model).score(
                execution,
                target_url=config.url,
                generated_tests_path=run_dir / "generated_tests.py",
            )
            run.scored_failures = scored
            for failure in scored:
                severity = failure.severity.upper()
                if severity in run.severity_breakdown:
                    run.severity_breakdown[severity] += 1
        except Exception:
            pass

    run.finished_at = datetime.now(UTC)
    HTMLReporter().generate(run)
    JSONReporter().generate(run)
    xlsx_path = build_xlsx(run, run_dir / "autonomous-qa-report.xlsx", srs_name, pen_name)
    (run_dir / "run_context.txt").write_text(context)
    return xlsx_path
