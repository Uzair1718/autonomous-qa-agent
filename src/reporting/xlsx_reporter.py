"""Excel-first QA report writer for n8n runs."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.models import RunData


def _fit(ws: Any) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min(max(max(len(str(c.value or "")) for c in col) + 2, 12), 60)
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"


def _header(ws: Any) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(bold=True, color="FFFFFF")


def build_xlsx(run: RunData, path: Path, srs_name: str = "", pen_name: str = "", visual_results: list[dict[str, Any]] | None = None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    rows = [
        ["Metric", "Value"],
        ["Run ID", run.run_id],
        ["Website", run.config.url],
        ["Pages Crawled", run.crawl_result.total_pages if run.crawl_result else 0],
        ["Flows Inferred", len(run.flows)],
        ["Tests Run", run.execution_result.total if run.execution_result else 0],
        ["Passed", run.execution_result.passed if run.execution_result else 0],
        ["Failed", run.execution_result.failed if run.execution_result else 0],
        ["Skipped", run.execution_result.skipped if run.execution_result else 0],
        ["Pass Rate", f"{run.pass_rate:.1f}%"],
        ["Critical", run.severity_breakdown.get("CRITICAL", 0)],
        ["High", run.severity_breakdown.get("HIGH", 0)],
        ["Medium", run.severity_breakdown.get("MEDIUM", 0)],
        ["Low", run.severity_breakdown.get("LOW", 0)],
        ["SRS", srs_name],
        ["Pen Design", pen_name],
    ]
    for row in rows:
        ws.append(row)
    _header(ws)
    _fit(ws)

    tc = wb.create_sheet("Test Cases")
    tc.append(["Test Case", "Status", "Duration (s)", "Issue / Error", "Trace"])
    for test in (run.execution_result.tests if run.execution_result else []):
        tc.append([test.name, test.status.upper(), round(test.duration, 2), test.error_message, str(test.trace_path or "")])
    _header(tc)
    _fit(tc)

    bugs = wb.create_sheet("Bug Report")
    bugs.append(["Bug Name", "Issue", "Image", "Severity", "Reproduction Steps", "Recommended Fix"])
    scored = {x.test_name: x for x in run.scored_failures}
    for test in (run.execution_result.tests if run.execution_result else []):
        if test.status != "failed":
            continue
        sf = scored.get(test.name)
        image_path = ""
        if test.trace_path:
            candidate = test.trace_path.parent.parent / "screenshots"
            if candidate.exists():
                images = list(candidate.glob("*.png"))
                if images:
                    image_path = str(images[0])
        bugs.append([
            test.name,
            test.error_message or "Test failed.",
            image_path,
            sf.severity if sf else "UNTRIAGED",
            "\n".join(sf.reproduction_steps) if sf else "",
            sf.recommended_fix if sf else "",
        ])
        if image_path and Path(image_path).exists():
            try:
                img = XLImage(image_path)
                img.width = 420
                img.height = 240
                bugs.add_image(img, f"C{bugs.max_row}")
            except Exception:
                pass
    _header(bugs)
    _fit(bugs)

    flows = wb.create_sheet("Use Cases")
    flows.append(["Use Case", "Priority", "Description", "Expected Outcome"])
    for flow in run.flows:
        flows.append([flow.name, flow.priority, flow.description, flow.expected_outcome])
    _header(flows)
    _fit(flows)

    req = wb.create_sheet("Requirements")
    req.append(["Requirement Source", "Requirement / Extract", "Coverage"])
    req.append(["SRS", "SRS was supplied to the runner; extracted requirement text is included in runner artifacts.", "AI flow planning context"])
    req.append(["Pen Design", "Pen design was supplied to the runner; structural design summary is included in runner artifacts.", "AI flow planning context"])
    _header(req)
    _fit(req)

    visual = wb.create_sheet("Pen Visual Comparison")
    visual.append(["Screen", "URL", "Status", "Similarity %", "Changed Area %", "Design Screenshot", "Live Screenshot", "Diff Screenshot", "Reason"])
    for item in visual_results or []:
        visual.append([
            item.get("screen", ""),
            item.get("url", ""),
            item.get("status", ""),
            item.get("similarity", ""),
            item.get("changed_ratio", ""),
            item.get("design_screenshot", ""),
            item.get("live_screenshot", ""),
            item.get("diff_screenshot", ""),
            item.get("reason", ""),
        ])
    _header(visual)
    _fit(visual)

    evidence = wb.create_sheet("Evidence")
    evidence.append(["Artifact", "Path"])
    for p in sorted(run.run_dir.rglob("*")):
        if p.is_file():
            evidence.append([p.name, str(p)])
    _header(evidence)
    _fit(evidence)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
