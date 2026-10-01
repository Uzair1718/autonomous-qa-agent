"""Excel QA report exporter."""
from __future__ import annotations

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font
from src.models import RunData

class ExcelReporter:
    def generate(self, run_data: RunData) -> Path:
        out = run_data.run_dir / "execution-report.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Test Cases"
        headers = ["Test Name","Status","Duration (s)","Page URL","Error"]
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
        for t in (run_data.execution_result.tests if run_data.execution_result else []):
            ws.append([t.name,t.status,round(t.duration,3),t.page_url,t.error_message])
        bugs = wb.create_sheet("Bugs")
        bug_headers = ["Bug ID","Bug Name","Issue","Steps to Reproduce","Expected Result","Actual Result","Severity","Priority","Page","Browser","User Role","Screenshot","Video","Trace","Console Error","Status"]
        bugs.append(bug_headers)
        for c in bugs[1]:
            c.font = Font(bold=True)
        for i, f in enumerate(run_data.scored_failures, 1):
            original = next((t for t in (run_data.execution_result.tests if run_data.execution_result else []) if t.name == f.test_name), None)
            bugs.append([f"BUG-{i:03d}",f.test_name,f.reason,"\n".join(f.reproduction_steps),f"Expected behavior per test flow",original.error_message if original else f.original_error,f.severity,f.severity,"",run_data.config.browsers[0] if run_data.config.browsers else "","", "", "", str(original.trace_path) if original and original.trace_path else "","", "Open"])
        use = wb.create_sheet("Test Scenarios")
        use.append(["Scenario","Priority","Description","Expected Outcome"])
        for c in use[1]: c.font = Font(bold=True)
        for f in run_data.flows:
            use.append([f.name,f.priority,f.description,f.expected_outcome])
        summary = wb.create_sheet("Summary")
        summary.append(["Metric","Value"])
        summary["A1"].font = summary["B1"].font = Font(bold=True)
        summary.append(["URL",run_data.config.url])
        summary.append(["Pages",run_data.crawl_result.total_pages if run_data.crawl_result else 0])
        summary.append(["Flows",len(run_data.flows)])
        summary.append(["Tests",run_data.execution_result.total if run_data.execution_result else 0])
        summary.append(["Passed",run_data.execution_result.passed if run_data.execution_result else 0])
        summary.append(["Failed",run_data.execution_result.failed if run_data.execution_result else 0])
        summary.append(["Pass Rate",round(run_data.pass_rate,2)])
        for sheet in wb.worksheets:
            for col in sheet.columns:
                letter = col[0].column_letter
                sheet.column_dimensions[letter].width = min(max(max((len(str(c.value or "")) for c in col), default=10)+2, 12), 60)
        wb.save(out)
        return out
